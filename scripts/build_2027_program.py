#!/usr/bin/env python3
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl"]
# ///
"""Import the public IC2S2 2027 program from one editable Excel workbook.

The workbook is intentionally safe to commit: it contains placeholder program data,
not registration records or private organizer exports.

Usage:
    uv run scripts/build_2027_program.py --seed
    uv run scripts/build_2027_program.py
"""

import argparse
import json
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "scripts/input/IC2S2_2027_program.xlsx"
OUTPUT = ROOT / "src/lib/data/full-program/program-2027.json"
LEGACY = ROOT / "src/lib/data/full-program"

HEADERS = [
    "day",
    "date",
    "event_id",
    "event_time",
    "event_type",
    "event_title",
    "event_location",
    "item_kind",
    "session_title",
    "track",
    "room",
    "item_id",
    "item_time",
    "paper_title",
    "authors",
    "theme",
    "keywords",
    "abstract",
]

DATES_2027 = [
    "Monday, July 26 — Tutorial Day",
    "Tuesday, July 27",
    "Wednesday, July 28",
    "Thursday, July 29",
]

ADJECTIVES = [
    "Adaptive",
    "Collective",
    "Digital",
    "Emergent",
    "Latent",
    "Networked",
    "Participatory",
    "Synthetic",
]
SUBJECTS = [
    "Communities",
    "Conversations",
    "Institutions",
    "Neighborhoods",
    "Networks",
    "Platforms",
    "Publics",
    "Systems",
]
METHODS = [
    "Across Scales",
    "in Context",
    "Through Simulation",
    "Using Digital Traces",
    "with Causal Models",
    "with Machine Learning",
    "with Mixed Methods",
    "with Network Analysis",
]
FIRST_NAMES = ["Alex", "Avery", "Casey", "Jordan", "Morgan", "Riley", "Robin", "Taylor"]
LAST_NAMES = ["Example", "Fiction", "Placeholder", "Sample", "Test", "Trial"]


def placeholder_title(identifier: int) -> str:
    return (
        f"{ADJECTIVES[identifier % len(ADJECTIVES)]} "
        f"{SUBJECTS[(identifier * 3) % len(SUBJECTS)]} "
        f"{METHODS[(identifier * 5) % len(METHODS)]}"
    )


def placeholder_authors(identifier: int) -> str:
    count = 2 + identifier % 3
    names = []
    for index in range(count):
        first = FIRST_NAMES[(identifier + index * 3) % len(FIRST_NAMES)]
        last = LAST_NAMES[(identifier * 2 + index) % len(LAST_NAMES)]
        names.append(f"{first} {last}")
    return ", ".join(names)


def placeholder_abstract() -> str:
    return (
        "This is placeholder abstract text for the draft IC2S2 2027 program. "
        "Replace this cell with the final abstract when the presentation is confirmed."
    )


def safe_location(event_type: str, index: int) -> str:
    if event_type == "registration":
        return "Registration desk"
    if event_type in {"remarks", "keynote", "lightning"}:
        return "Main auditorium"
    if event_type == "poster":
        return "Poster hall"
    if event_type == "social":
        return "Reception area"
    return ""


def seed_workbook() -> None:
    """Create an editable 2027 workbook using the existing schedule shape."""
    program = json.loads((LEGACY / "program.json").read_text())
    sessions = {
        session["title"]: session
        for session in json.loads((LEGACY / "parallel-sessions.json").read_text())
    }
    posters_by_date = {
        session["date"]: session
        for session in json.loads((LEGACY / "poster.json").read_text())
    }

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Program"
    sheet.append(HEADERS)

    placeholder_id = 1
    for day_index, day in enumerate(program):
        for event_index, event in enumerate(day["events"]):
            event_id = f"D{day_index + 1}-E{event_index + 1}"
            event_type = event["type"]
            event_title = event["title"]
            if event_type == "keynote":
                event_title = "Keynote — speaker to be announced"
            elif event_type == "tutorial":
                event_title = "Tutorial details to be announced"

            base = {
                "day": day["day"],
                "date": DATES_2027[day_index],
                "event_id": event_id,
                "event_time": event["time"],
                "event_type": event_type,
                "event_title": event_title,
                "event_location": safe_location(event_type, event_index),
            }

            rows = []
            if event.get("items"):
                for item in event["items"]:
                    rows.append(
                        base
                        | {
                            "item_kind": "lightning",
                            "item_id": placeholder_id,
                            "item_time": item.get("time", ""),
                            "paper_title": placeholder_title(placeholder_id),
                            "authors": placeholder_authors(placeholder_id),
                            "abstract": placeholder_abstract(),
                        }
                    )
                    placeholder_id += 1
            elif event.get("parallelSessions"):
                for session_name in event["parallelSessions"]:
                    session = sessions[session_name]
                    for paper in session["papers"]:
                        rows.append(
                            base
                            | {
                                "item_kind": "parallel",
                                "session_title": session_name,
                                "track": session.get("track", ""),
                                "room": f"Room {session.get('track', '')}",
                                "item_id": placeholder_id,
                                "paper_title": placeholder_title(placeholder_id),
                                "authors": placeholder_authors(placeholder_id),
                                "abstract": placeholder_abstract(),
                            }
                        )
                        placeholder_id += 1
            elif event_type == "poster":
                poster_session = posters_by_date.get(day["date"])
                for poster in (poster_session or {}).get("posters", []):
                    rows.append(
                        base
                        | {
                            "item_kind": "poster",
                            "session_title": event_title,
                            "item_id": placeholder_id,
                            "paper_title": placeholder_title(placeholder_id),
                            "authors": placeholder_authors(placeholder_id),
                            "theme": "Placeholder theme",
                            "keywords": "placeholder; draft; example",
                            "abstract": placeholder_abstract(),
                        }
                    )
                    placeholder_id += 1
            else:
                rows.append(base | {"item_kind": "event"})

            for row in rows:
                sheet.append([row.get(header, "") for header in HEADERS])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="484D55")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [12, 34, 12, 16, 16, 38, 22, 14, 46, 10, 14, 10, 12, 62, 42, 24, 32, 80]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK)
    print(f"Wrote placeholder workbook: {WORKBOOK.relative_to(ROOT)}")


def clean(value) -> str:
    return str(value or "").strip()


def import_workbook() -> None:
    """Convert the public workbook to the nested JSON consumed by Svelte."""
    if not WORKBOOK.exists():
        raise SystemExit(f"Missing {WORKBOOK}. Run this script with --seed first.")

    sheet = load_workbook(WORKBOOK, read_only=True, data_only=True)["Program"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [clean(value) for value in rows[0]]
    missing_headers = [header for header in HEADERS if header not in headers]
    if missing_headers:
        raise SystemExit(
            "The Program worksheet is missing required columns: "
            + ", ".join(missing_headers)
            + ". Restore the original column names in row 1."
        )

    # Excel and other spreadsheet editors commonly omit trailing empty cells when
    # saving. Map every declared header explicitly so event-only rows still receive
    # blank values for columns such as item_id, authors, and abstract.
    records = [
        {
            header: row[index] if index < len(row) else None
            for index, header in enumerate(headers)
        }
        for row in rows[1:]
        if any(row)
    ]

    days: OrderedDict[str, dict] = OrderedDict()
    event_maps: dict[str, OrderedDict[str, dict]] = {}

    for record in records:
        day_name = clean(record["day"])
        event_id = clean(record["event_id"])
        if not day_name or not event_id:
            raise SystemExit("Every row must have day and event_id values.")

        if day_name not in days:
            days[day_name] = {
                "day": day_name,
                "date": clean(record["date"]),
                "events": [],
            }
            event_maps[day_name] = OrderedDict()

        events = event_maps[day_name]
        if event_id not in events:
            event = {
                "time": clean(record["event_time"]),
                "title": clean(record["event_title"]),
                "type": clean(record["event_type"]),
            }
            if clean(record["event_location"]):
                event["location"] = clean(record["event_location"])
            events[event_id] = event
            days[day_name]["events"].append(event)

        event = events[event_id]
        kind = clean(record["item_kind"])
        item_id = 0
        if kind in {"lightning", "parallel", "poster"}:
            raw_item_id = record["item_id"]
            if raw_item_id in (None, ""):
                raise SystemExit(
                    f"Missing item_id for {kind} row in event {event_id}. "
                    "Every paper, poster, and lightning talk needs a numeric item_id."
                )
            try:
                item_id = int(float(raw_item_id))
            except (TypeError, ValueError):
                raise SystemExit(
                    f"Invalid item_id {raw_item_id!r} in event {event_id}; expected a number."
                ) from None

        if kind == "lightning":
            event.setdefault("items", []).append(
                {
                    "time": clean(record["item_time"]),
                    "title": clean(record["paper_title"]),
                    "presenters": clean(record["authors"]),
                    "abstract": clean(record["abstract"]),
                }
            )
        elif kind == "parallel":
            sessions = event.setdefault("parallelSessions", [])
            session_title = clean(record["session_title"])
            session = next((item for item in sessions if item["title"] == session_title), None)
            if session is None:
                session = {
                    "title": session_title,
                    "day": len(days),
                    "time": "AM" if clean(record["event_time"]).split(":")[0] < "12" else "PM",
                    "track": clean(record["track"]),
                    "room": clean(record["room"]),
                    "papers": [],
                }
                sessions.append(session)
            session["papers"].append(
                {
                    "submission": item_id,
                    "title": clean(record["paper_title"]),
                    "authors": clean(record["authors"]),
                    "abstract": clean(record["abstract"]),
                }
            )
        elif kind == "poster":
            poster_session = event.setdefault(
                "posterSession",
                {
                    "day": len(days),
                    "session": clean(record["session_title"]),
                    "date": days[day_name]["date"],
                    "posters": [],
                },
            )
            poster_session["posters"].append(
                {
                    "id": item_id,
                    "title": clean(record["paper_title"]),
                    "authors": clean(record["authors"]),
                    "theme": clean(record["theme"]),
                    "keywords": clean(record["keywords"]),
                    "abstract": clean(record["abstract"]),
                }
            )

    OUTPUT.write_text(json.dumps(list(days.values()), ensure_ascii=False, indent=2) + "\n")
    print(f"Imported {len(records)} spreadsheet rows to {OUTPUT.relative_to(ROOT)}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--seed", action="store_true", help="create the initial placeholder workbook")
    args = parser.parse_args()
    if args.seed:
        seed_workbook()
    import_workbook()
